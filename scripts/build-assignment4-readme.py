from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "WebClient" / "README.md"
OUTPUT = ROOT / "target" / "assignment4" / "Guess_Market_Ex4_README.docx"
TOKENS = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|https?://\S+)")


def bidi(paragraph):
    element = OxmlElement("w:bidi")
    element.set(qn("w:val"), "1")
    paragraph._p.get_or_add_pPr().append(element)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def add_inline(paragraph, content):
    for token in TOKENS.split(content):
        if not token:
            continue
        bold = token.startswith("**")
        code = token.startswith("`")
        value = token[2:-2] if bold else token[1:-1] if code else token
        run = paragraph.add_run(value)
        run.bold = bold
        run.font.rtl = not (code or value.startswith("http"))
        if code or value.startswith("http"):
            run.font.name = "Consolas"
            run.font.size = Pt(9)


document = Document()
section = document.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Cm(2)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2)
section.right_margin = Cm(2)

normal = document.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(30, 38, 42)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.18

title = document.styles["Title"]
title.font.name = "Arial"
title.font.size = Pt(18)
title.font.bold = True
title.font.color.rgb = RGBColor(20, 28, 32)
title.paragraph_format.space_after = Pt(11)
title_border = title._element.get_or_add_pPr().find(qn("w:pBdr"))
if title_border is not None:
    title._element.pPr.remove(title_border)

heading = document.styles["Heading 1"]
heading.font.name = "Arial"
heading.font.size = Pt(13)
heading.font.bold = True
heading.font.color.rgb = RGBColor(15, 73, 85)
heading.paragraph_format.space_before = Pt(15)
heading.paragraph_format.space_after = Pt(6)
heading.paragraph_format.keep_with_next = True

for raw in SOURCE.read_text(encoding="utf-8").splitlines():
    line = raw.strip()
    if not line:
        continue
    if line.startswith("# "):
        paragraph = document.add_paragraph(style="Title")
        value = line[2:]
    elif line.startswith("## "):
        paragraph = document.add_paragraph(style="Heading 1")
        value = line[3:]
    elif line.startswith("- "):
        paragraph = document.add_paragraph()
        value = "• " + line[2:]
        paragraph.paragraph_format.right_indent = Cm(0.25)
        paragraph.paragraph_format.space_after = Pt(4)
    elif re.match(r"^\d+\. ", line):
        paragraph = document.add_paragraph()
        value = line
        paragraph.paragraph_format.right_indent = Cm(0.25)
        paragraph.paragraph_format.space_after = Pt(6)
    else:
        paragraph = document.add_paragraph()
        value = line
    bidi(paragraph)
    add_inline(paragraph, value)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(OUTPUT)
